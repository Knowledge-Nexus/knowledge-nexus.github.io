"""Geradores de ficheiros FICTÍCIOS para os testes (nada binário vai para o git).

Todos os textos são inventados. Uso:
    from amostras.gerar import pdf_nativo, pdf_digitalizado, ...
ou, para gerar uma pasta de exemplo:
    uv run python tests/amostras/gerar.py <destino>
"""

from __future__ import annotations

import io
import json
import sys
import zipfile
from pathlib import Path

FONT_PATHS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
]

EXAME_AM1 = """Universidade Fictícia de Exemplo
Licenciatura em Engenharia Informática
Unidade Curricular: Análise Matemática I
Exame - Época de Recurso
Ano Lectivo 2023/2024 - 5 de Fevereiro de 2024

Grupo I (6 valores)
1. Calcule o limite da sucessão u_n = (2n + 1) / (n + 3).
2. Estude a continuidade da função f(x) = |x| / x no ponto zero.

Grupo II (8 valores)
3. Calcule a derivada de g(x) = x^2 sin(x).
4. Determine uma primitiva de h(x) = 1 / (1 + x^2).
"""

RESOLUCAO_AM1 = """Universidade Fictícia de Exemplo
Análise Matemática I - Proposta de Resolução
Exame - Época de Recurso - Ano Lectivo 2023/2024

1. O limite da sucessão é 2, porque o quociente dos termos de maior grau é 2n/n.
2. A função não está definida em zero, logo não é contínua nesse ponto.
3. g'(x) = 2x sin(x) + x^2 cos(x).
4. Uma primitiva é arctg(x).
"""

SLIDES_ALGA = """Álgebra Linear e Geometria Analítica
Aula teórica 3 - Matrizes e determinantes
Docente: Carla Modelo

Uma matriz quadrada é invertível se e só se o determinante for diferente de zero.
O determinante de uma matriz triangular é o produto dos elementos da diagonal.
"""

FICHA_DESCONHECIDA = """Instituto Superior de Exemplo
Unidade Curricular: Teoria dos Grafos Imaginários
Ficha de exercícios 2

1. Mostre que todo o grafo finito tem um número par de vértices de grau ímpar.
2. Determine uma árvore abrangente do grafo da figura.
"""

APONTAMENTOS_FG = """Física Geral - apontamentos das aulas
Cinemática: a velocidade é a derivada da posição em ordem ao tempo.
Dinâmica: a segunda lei de Newton relaciona a força com a aceleração.
Energia cinética e energia potencial; conservação da energia mecânica.
"""


def _font(size: int):  # type: ignore[no-untyped-def]
    from PIL import ImageFont

    for candidate in FONT_PATHS:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def pdf_nativo(path: Path, text: str, title: str | None = None) -> Path:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # Com uma fonte TrueType cabem todos os caracteres (ex.: os acentos soltos do LaTeX);
    # sem ela, fica a Helvetica de base.
    fontfile = next((f for f in FONT_PATHS if Path(f).exists()), None)
    font = {"fontname": "dejavu", "fontfile": fontfile} if fontfile else {"fontname": "helv"}
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11, **font)
    if title:
        doc.set_metadata({"title": title})
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    return path


def imagem_texto(path: Path, text: str, rotate: int = 0) -> Path:
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    font = _font(40)
    y = 120
    for line in text.splitlines():
        draw.text((120, y), line, fill="black", font=font)
        y += 62
    if rotate:
        image = image.rotate(rotate, expand=True, fillcolor="white")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)
    return path


def pdf_digitalizado(path: Path, text: str, workdir: Path) -> Path:
    """PDF só com imagem (sem camada de texto), como um documento digitalizado."""
    import pymupdf

    png = imagem_texto(workdir / f"{path.stem}-scan.png", text)
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, filename=str(png))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()
    png.unlink()
    return path


def docx_com_equacao(path: Path, text: str) -> Path:
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    for index, line in enumerate(text.splitlines()):
        if not line.strip():
            continue
        if index == 0:
            document.add_heading(line, level=1)
        else:
            document.add_paragraph(line)
    # Equação OMML (E = m c²), para verificar a conversão para LaTeX pelo Pandoc.
    omml = (
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        "<m:oMath><m:r><m:t>E=m</m:t></m:r><m:sSup><m:e><m:r><m:t>c</m:t></m:r></m:e>"
        "<m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup></m:oMath></m:oMathPara>"
    )
    document.add_paragraph()._p.append(parse_xml(omml))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def pptx_slides(path: Path, title: str, bullets: list[str]) -> Path:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for bullet in bullets[1:]:
        body.add_paragraph().text = bullet
    slide.notes_slide.notes_text_frame.text = "Notas fictícias do docente."
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path


def xlsx_notas(path: Path) -> Path:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Pautas"
    sheet.append(["Número", "Nome", "Nota"])
    sheet.append([1001, "Aluno Fictício", 14])
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def projecto_codigo(folder: Path) -> Path:
    (folder / "src").mkdir(parents=True, exist_ok=True)
    (folder / "pyproject.toml").write_text('[project]\nname = "trabalho-p1"\n')
    (folder / "src" / "main.py").write_text("def soma(a, b):\n    return a + b\n")
    (folder / "src" / "util.py").write_text("PI = 3.14159\n")
    (folder / "README.md").write_text("# Trabalho prático de Programação I (fictício)\n")
    (folder / "node_modules" / "lib").mkdir(parents=True, exist_ok=True)
    (folder / "node_modules" / "lib" / "index.js").write_text("module.exports = 1;\n")
    return folder


def zip_de(path: Path, members: dict[str, bytes]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    return path


def pdf_bytes(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # Com uma fonte TrueType cabem todos os caracteres (ex.: os acentos soltos do LaTeX);
    # sem ela, fica a Helvetica de base.
    fontfile = next((f for f in FONT_PATHS if Path(f).exists()), None)
    font = {"fontname": "dejavu", "fontfile": fontfile} if fontfile else {"fontname": "helv"}
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11, **font)
    buffer = io.BytesIO()
    doc.save(buffer)
    doc.close()
    return buffer.getvalue()


def notebook(path: Path) -> Path:
    data = {
        "cells": [
            {"cell_type": "markdown", "source": ["# Ficha 1 de Programação I\n"]},
            {"cell_type": "code", "source": ["print('olá')\n"]},
        ],
        "metadata": {"kernelspec": {"language": "python"}},
    }
    path.write_text(json.dumps(data))
    return path


def pasta_exemplo(dest: Path) -> Path:
    """Um lote de depósito variado, para experimentar o pipeline à mão."""
    lote = dest / "exemplo"
    pdf_nativo(lote / "AM1" / "Exame_Recurso_2023-24.pdf", EXAME_AM1)
    pdf_nativo(lote / "AM1" / "Exame_Recurso_2023-24_resolucao.pdf", RESOLUCAO_AM1)
    pptx_slides(lote / "ALGA_T3_matrizes.pptx", "Matrizes e determinantes",
                ["Matriz invertível", "Determinante"])
    docx_com_equacao(lote / "apontamentos_FG.docx", APONTAMENTOS_FG)
    pdf_nativo(lote / "grafos_ficha2.pdf", FICHA_DESCONHECIDA)
    projecto_codigo(lote / "trabalho-p1")
    return lote


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "amostras-geradas")
    print(pasta_exemplo(target))
