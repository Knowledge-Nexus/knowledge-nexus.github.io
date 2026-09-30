"""Documentos Office: DOCX (Pandoc → Markdown com LaTeX), PPTX, XLSX e formatos antigos
(.doc, .ppt, .xls, .odt… convertidos pelo LibreOffice)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from nexus.config import ExtractionSettings
from nexus.pipeline.extract.base import (
    ExtractionError,
    ExtractionResult,
    PageText,
    clean_text,
    require,
    which,
)
from nexus.pipeline.extract.pdf import pdf_pages

DOCX_VERSION = 1
PPTX_VERSION = 1
XLSX_VERSION = 1
LEGACY_VERSION = 1


def libreoffice_convert(path: Path, target: str, workdir: Path, timeout: int) -> Path:
    """Converte com o LibreOffice sem interface; devolve o ficheiro gerado."""
    soffice = require("soffice", "conversão de documentos Office")
    outdir = workdir / f"lo-{target}"
    outdir.mkdir(parents=True, exist_ok=True)
    profile = (workdir / "lo-profile").resolve().as_uri()
    source = workdir / f"input{path.suffix.lower()}"
    if source.resolve() != path.resolve():
        shutil.copyfile(path, source)
    try:
        proc = subprocess.run(
            [soffice, "--headless", "--norestore", "--nolockcheck",
             f"-env:UserInstallation={profile}", "--convert-to", target,
             "--outdir", str(outdir), str(source)],
            capture_output=True, text=True, timeout=timeout, check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise ExtractionError(f"o LibreOffice excedeu {timeout}s") from exc
    produced = outdir / f"{source.stem}.{target.split(':', 1)[0]}"
    if proc.returncode != 0 or not produced.exists():
        raise ExtractionError(f"conversão para {target} falhou: {proc.stderr.strip()[:300]}")
    return produced


def _rendition(
    path: Path, settings: ExtractionSettings, workdir: Path, warnings: list[str]
) -> Path | None:
    """Versão PDF (para abrir na página certa). É opcional: se falhar, segue sem ela."""
    if not which("soffice"):
        warnings.append("LibreOffice indisponível: sem paginação nem versão PDF")
        return None
    try:
        return libreoffice_convert(path, "pdf", workdir, settings.office_timeout_seconds)
    except ExtractionError as exc:
        warnings.append(f"sem versão PDF: {exc}")
        return None


def _pandoc_markdown(path: Path, timeout: int) -> str | None:
    if not which("pandoc"):
        return None
    fmt = ("markdown-raw_html-native_divs-native_spans-fenced_divs-bracketed_spans"
           "-header_attributes-link_attributes-simple_tables-multiline_tables-grid_tables")
    proc = subprocess.run(
        ["pandoc", str(path), "-f", "docx", "-t", fmt, "--wrap=none"],
        capture_output=True, text=True, timeout=timeout, check=False,
    )
    if proc.returncode != 0:
        return None
    return clean_text(proc.stdout)


def _python_docx_markdown(path: Path) -> str:
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"DOCX ilegível: {exc}") from exc
    lines: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading") or style.startswith("título"):
            level = next((c for c in style if c.isdigit()), "1")
            lines.append(f"{'#' * int(level)} {text}")
        else:
            lines.append(text)
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        lines.append(_markdown_table(rows))
    return clean_text("\n\n".join(lines))


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [[c.replace("|", "\\|").replace("\n", " ") for c in r] + [""] * (width - len(r))
            for r in rows]
    out = ["| " + " | ".join(norm[0]) + " |", "|" + " --- |" * width]
    out += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(out)


def extract_docx(path: Path, settings: ExtractionSettings, workdir: Path) -> ExtractionResult:
    warnings: list[str] = []
    markdown = _pandoc_markdown(path, settings.office_timeout_seconds)
    if markdown is None:
        warnings.append("Pandoc indisponível: equações podem não ter sido convertidas para LaTeX")
        markdown = _python_docx_markdown(path)
    rendition = _rendition(path, settings, workdir, warnings)
    pages = pdf_pages(rendition, settings, workdir, allow_ocr=False)[0] if rendition else []
    if not pages or not any(p.text for p in pages):
        pages = [PageText(markdown)]
    return ExtractionResult("docx", DOCX_VERSION, pages, {}, markdown, rendition, warnings)


def extract_pptx(path: Path, settings: ExtractionSettings, workdir: Path) -> ExtractionResult:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(f"PPTX ilegível: {exc}") from exc
    pages: list[PageText] = []
    for slide in presentation.slides:
        title = ""
        blocks: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                is_title = shape.is_placeholder and "title" in str(
                    shape.placeholder_format.type).lower()
                text = "\n".join(
                    ("  " * p.level + "- " if not is_title and p.level else "")
                    + "".join(run.text for run in p.runs)
                    for p in shape.text_frame.paragraphs
                ).strip()
                if not text:
                    continue
                if is_title and not title:
                    title = text.replace("\n", " ")
                else:
                    blocks.append(text)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                blocks.append(_markdown_table(rows))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append("> **Notas:** " + notes.replace("\n", " "))
        body = "\n\n".join(([f"## {title}"] if title else []) + blocks)
        pages.append(PageText(clean_text(body), label=title or None))
    warnings: list[str] = []
    rendition = _rendition(path, settings, workdir, warnings)
    return ExtractionResult("pptx", PPTX_VERSION, pages, {}, None, rendition, warnings)


def extract_xlsx(path: Path, settings: ExtractionSettings, workdir: Path) -> ExtractionResult:
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"folha de cálculo ilegível: {exc}") from exc
    pages: list[PageText] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(max_row=settings.max_xlsx_rows,
                                   max_col=settings.max_xlsx_cols, values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                rows.append(cells)
        body = f"## {sheet.title}\n\n{_markdown_table(rows)}" if rows else f"## {sheet.title}"
        pages.append(PageText(body, label=sheet.title))
    workbook.close()
    return ExtractionResult("xlsx", XLSX_VERSION, pages)


def extract_legacy(
    path: Path, target_ext: str, settings: ExtractionSettings, workdir: Path
) -> ExtractionResult:
    converted = libreoffice_convert(path, target_ext, workdir, settings.office_timeout_seconds)
    handler = {"docx": extract_docx, "pptx": extract_pptx, "xlsx": extract_xlsx}[target_ext]
    result = handler(converted, settings, workdir)
    result.extractor = f"legacy-{target_ext}"
    result.version = LEGACY_VERSION * 100 + result.version
    return result
